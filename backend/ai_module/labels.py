"""
Связывает экспертную разметку (`Разметка.xlsx`, лист «Калибровка») со снимками
обучающей выборки (шаг 0 плана переделки, `docs/ПЛАН_ПЕРЕДЕЛКИ.md`).

Идентификатор `study` в xlsx буквально совпадает с именем папки исследования в
`НД_для_обучения/Исследования/<study>/...` — проверено по всем 100 строкам
(100/100 точных строковых совпадений). Отдельная логика сопоставления через
DICOM-теги (StudyInstanceUID и т.п.) не потребовалась.

Сырой датасет не хранится в репозитории (персональные медицинские данные,
см. `.gitignore`) — пути к xlsx и к папке с исследованиями передаются вызывающим
кодом явно.
"""
from __future__ import annotations

from pathlib import Path
from typing import TypedDict

import openpyxl

SHEET_NAME = "Калибровка"
FIRST_DATA_ROW = 3
LAST_DATA_ROW = 102

# Столбцы B..M листа «Калибровка» (в этом порядке идут в строке).
COL_STUDY = 0
COL_SPINE_POSITIONING = 1     # C
COL_SPINE_AXIS = 2            # D
COL_SPINE_ARTIFACT = 3        # E
COL_HIP_R_POSITIONING = 4     # F
COL_HIP_R_FOV = 5             # G
COL_HIP_L_POSITIONING = 6     # H
COL_HIP_L_FOV = 7             # I
COL_SPINE_OVERALL = 8         # J
COL_HIP_R_OVERALL = 9         # K
COL_HIP_L_OVERALL = 10        # L
COL_COMMENT = 11              # M


class RegionLabel(TypedDict):
    incorrect_positioning: bool | None
    overall: bool | None


class SpineLabel(RegionLabel):
    axis_misaligned: bool | None
    artifact: bool | None


class HipLabel(RegionLabel):
    incomplete_field_of_view: bool | None


class StudyLabel(TypedDict):
    spine: SpineLabel
    hip_right: HipLabel
    hip_left: HipLabel
    comment: str | None


def _as_bool(value) -> bool | None:
    """0/1 из xlsx -> bool, пустая ячейка (область не размечена) -> None."""
    if value is None or value == "":
        return None
    return bool(int(value))


def load_labels(xlsx_path: Path | str) -> dict[str, StudyLabel]:
    """Читает `Разметка.xlsx` и возвращает {study_dir: StudyLabel} по всем строкам."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]

    labels: dict[str, StudyLabel] = {}
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        values = [ws.cell(row=row, column=c).value for c in range(2, 14)]  # B..M
        study = values[COL_STUDY]
        if not study:
            continue
        labels[str(study).strip()] = {
            "spine": {
                "incorrect_positioning": _as_bool(values[COL_SPINE_POSITIONING]),
                "axis_misaligned": _as_bool(values[COL_SPINE_AXIS]),
                "artifact": _as_bool(values[COL_SPINE_ARTIFACT]),
                "overall": _as_bool(values[COL_SPINE_OVERALL]),
            },
            "hip_right": {
                "incorrect_positioning": _as_bool(values[COL_HIP_R_POSITIONING]),
                "incomplete_field_of_view": _as_bool(values[COL_HIP_R_FOV]),
                "overall": _as_bool(values[COL_HIP_R_OVERALL]),
            },
            "hip_left": {
                "incorrect_positioning": _as_bool(values[COL_HIP_L_POSITIONING]),
                "incomplete_field_of_view": _as_bool(values[COL_HIP_L_FOV]),
                "overall": _as_bool(values[COL_HIP_L_OVERALL]),
            },
            "comment": values[COL_COMMENT],
        }
    return labels


def link_labels_to_studies(
    xlsx_path: Path | str, studies_dir: Path | str
) -> tuple[dict[str, StudyLabel], list[str]]:
    """Связывает разметку с реально существующими папками исследований.

    Возвращает (labels_by_existing_dir, missing) — `missing` содержит study id
    из xlsx, для которых не нашлось папки `studies_dir/<study>`.
    """
    labels = load_labels(xlsx_path)
    studies_dir = Path(studies_dir)
    existing = {p.name for p in studies_dir.iterdir() if p.is_dir()}

    linked: dict[str, StudyLabel] = {}
    missing: list[str] = []
    for study, label in labels.items():
        if study in existing:
            linked[study] = label
        else:
            missing.append(study)
    return linked, missing


if __name__ == "__main__":
    import sys

    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("разметка.xlsx")
    studies = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("Исследования")

    linked, missing = link_labels_to_studies(xlsx, studies)
    print(f"связано: {len(linked)} из {len(linked) + len(missing)}")
    if missing:
        print("без папки:", ", ".join(missing))
