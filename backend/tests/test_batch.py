"""Пакетная обработка (ai_module/batch.py): папка / zip / файл -> таблица ТЗ п.2.5,
и REST API POST /batch. Данные — синтетические DICOM из sample_data/studies."""
import csv
import io
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session, SQLModel, create_engine
from sqlalchemy.pool import StaticPool

from ai_module.batch import COLUMNS, extract_zip, find_dicom_files, run_batch, table_bytes

SAMPLES = Path(__file__).resolve().parent.parent.parent / "sample_data" / "studies"
OK_SPINE = SAMPLES / "01_lspine_ok.dcm"
OK_HIP = SAMPLES / "02_hip_ok.dcm"


class _Cp866Info(zipfile.ZipInfo):
    """Имя пишется байтами cp866 без флага UTF-8 — как у русских архиваторов
    Windows (сам zipfile не-ASCII имена всегда пишет в UTF-8 с флагом)."""

    def _encodeFilenameFlags(self):
        return self.filename.encode("cp866"), self.flag_bits & ~0x800


def _cp866_entry(name: str) -> zipfile.ZipInfo:
    return _Cp866Info(name)


def _make_archive(path: Path) -> None:
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as zi:
        zi.write(OK_HIP, "hip/CR000002")                       # DICOM без расширения
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(_cp866_entry("Исследования/пациент_1/CR000001.dcm"), OK_SPINE.read_bytes())
        zf.writestr("nested.zip", inner.getvalue())             # вложенный архив
        zf.writestr("../escape.dcm", OK_SPINE.read_bytes())     # попытка выйти из папки
        zf.writestr("readme.txt", "не DICOM")
        zf.writestr("__MACOSX/._CR000001.dcm", b"junk")


def test_extract_zip_names_nesting_and_traversal(tmp_path):
    archive = tmp_path / "set.zip"
    _make_archive(archive)
    out = tmp_path / "out"
    extract_zip(archive, out)

    found = sorted(p.relative_to(out).as_posix() for p in find_dicom_files(out))
    assert found == sorted(["Исследования/пациент_1/CR000001.dcm", "escape.dcm", "nested/hip/CR000002"])
    assert not (tmp_path / "escape.dcm").exists()   # не вышло за пределы out
    assert not (out / "nested.zip").exists()        # вложенный архив раскрыт


def test_run_batch_zip(tmp_path):
    archive = tmp_path / "set.zip"
    _make_archive(archive)
    rows = run_batch(archive, "heuristics")
    by_path = {r["path_to_study"]: r for r in rows}
    assert set(by_path) == {"set.zip/Исследования/пациент_1/CR000001.dcm",
                            "set.zip/escape.dcm", "set.zip/nested/hip/CR000002"}
    assert all(r["processing_status"] == "Success" for r in rows)
    assert by_path["set.zip/nested/hip/CR000002"]["anatomical_region"] == "Проксимальный отдел бедра"
    for r in rows:
        assert r["quality_class"] in (0, 1) and 0.0 <= r["quality_prob"] <= 1.0


def test_run_batch_dir_and_single_file(tmp_path):
    shutil.copy(OK_SPINE, tmp_path / "a.dcm")
    (tmp_path / "notes.txt").write_text("x")
    assert [Path(r["path_to_study"]).name for r in run_batch(tmp_path, "heuristics")] == ["a.dcm"]
    single = run_batch(OK_SPINE, "heuristics")
    assert len(single) == 1 and single[0]["anatomical_region"] == "Поясничный отдел позвоночника"


def test_broken_dcm_is_failure_not_crash(tmp_path):
    (tmp_path / "broken.dcm").write_bytes(b"not a dicom at all")
    rows = run_batch(tmp_path, "heuristics")
    assert len(rows) == 1
    assert rows[0]["processing_status"] == "Failure"
    assert rows[0]["quality_class"] == "" and rows[0]["violation_type"] == ""
    assert rows[0]["error"]


def test_table_bytes_formats():
    rows = run_batch(OK_SPINE, "heuristics")
    text = table_bytes(rows, "csv").decode("utf-8-sig")
    header = next(csv.reader(io.StringIO(text)))
    assert header == COLUMNS                          # служебное поле error в таблицу не попадает
    ws = openpyxl.load_workbook(io.BytesIO(table_bytes(rows, "xlsx"))).active
    assert [c.value for c in ws[1]] == COLUMNS and ws.max_row == 2


@pytest.fixture
def api_client():
    """Приложение с временной in-memory БД и подставленным пользователем:
    рабочая БД backend/data не трогается (startup-события без `with` не запускаются)."""
    from app.database import get_session
    from app.main import app
    from app.models import Role, User
    from app.security import get_current_user

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)
    with Session(engine) as s:
        user = User(username="tester", full_name="Тест", role=Role.expert, password_hash="x")
        s.add(user)
        s.commit()
        s.refresh(user)

    def _session():
        with Session(engine) as s:
            yield s

    app.dependency_overrides[get_session] = _session
    app.dependency_overrides[get_current_user] = lambda: user
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_api_batch_zip_csv(api_client, tmp_path):
    archive = tmp_path / "set.zip"
    _make_archive(archive)
    with open(archive, "rb") as f:
        r = api_client.post("/batch?format=csv", files={"file": ("set.zip", f, "application/zip")})
    assert r.status_code == 200, r.text
    assert r.headers["x-processed-files"] == "3" and r.headers["x-failed-files"] == "0"
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert len(rows) == 3 and list(rows[0]) == COLUMNS


def test_api_batch_xlsx_single_dicom(api_client):
    with open(OK_HIP, "rb") as f:
        r = api_client.post("/batch", files={"file": ("02_hip_ok.dcm", f, "application/dicom")})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2 and ws["A2"].value == "02_hip_ok.dcm"


def test_api_batch_rejects_non_dicom(api_client):
    r = api_client.post("/batch", files={"file": ("x.txt", b"hello", "text/plain")})
    assert r.status_code == 400
